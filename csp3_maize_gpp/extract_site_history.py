"""Parse 2001-2017 CSP3 site-history plant/harvest dates.

Uses sheet CSP3 only. First PLANT (skipping replants) to last HARVEST, then ±14 days.
Crop comes from notes/titles/hybrids — not from an odd/even rotation assumption.
"""

from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from config import (
    BUFFER_DAYS,
    LOCAL_WORKBOOKS_DIR,
    OPS_CHECKPOINTS,
    SITE_HISTORY_FILE,
    CheckpointError,
)
from formulas import buffer_window, gee_end_exclusive
from headers import normalize_header

SKIP_PLANT_RE = re.compile(
    r"replant|flooded|complete planting|preplant|fill[- ]?in",
    re.I,
)

MAIZE_WORD_RE = re.compile(r"\b(maize|corn)\b", re.I)
SOY_WORD_RE = re.compile(r"\b(soybeans?|soy)\b", re.I)
PIONEER_MAIZE_RE = re.compile(r"pioneer\s*(33\w+|14\d{2}\w*)", re.I)
PIONEER_SOY_RE = re.compile(r"pioneer\s*93", re.I)
DEKALB_RE = re.compile(r"\bdekalb\b|\bdkb?\b", re.I)
ASGROW_RE = re.compile(r"\basgrow\b", re.I)


def _as_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _text(*parts) -> str:
    bits = []
    for part in parts:
        if part is None:
            continue
        bits.append(str(part).replace("\u200b", " ").strip())
    return " ".join(b for b in bits if b)


def infer_crop(*texts: str) -> tuple[str | None, str]:
    blob = _text(*texts)
    if SOY_WORD_RE.search(blob):
        return "soybean", "notes_or_title"
    if MAIZE_WORD_RE.search(blob):
        return "maize", "notes_or_title"
    if PIONEER_SOY_RE.search(blob) or ASGROW_RE.search(blob):
        return "soybean", "hybrid_code"
    if PIONEER_MAIZE_RE.search(blob) or DEKALB_RE.search(blob):
        return "maize", "hybrid_code"
    return None, "unresolved"


def should_skip_plant(title: str | None, notes: str | None) -> bool:
    blob = _text(title, notes)
    return bool(SKIP_PLANT_RE.search(blob))


def _parse_structured(ws) -> list[dict]:
    events = []
    for row in range(5, 130):
        year = ws.cell(row, 1).value
        activity = ws.cell(row, 4).value
        if not activity:
            continue
        act = str(activity).strip().upper()
        if act not in {"PLANT", "HARVEST"}:
            continue
        event_date = _as_date(ws.cell(row, 3).value)
        notes = ws.cell(row, 6).value
        events.append(
            {
                "section": "structured",
                "year": int(year) if year else (event_date.year if event_date else None),
                "date": event_date,
                "activity": "PLANT" if act == "PLANT" else "HARVEST",
                "title": None,
                "notes": notes,
                "cell": f"C{row}",
                "activity_cell": f"D{row}",
                "notes_cell": f"F{row}",
            }
        )
    return events


def _parse_appendix(ws) -> list[dict]:
    events = []
    max_row = ws.max_row or 1
    for row in range(130, max_row + 1):
        stamp = ws.cell(row, 2).value
        event_date = _as_date(stamp)
        if event_date is None:
            continue
        title = ws.cell(row, 3).value
        next_kind = ws.cell(row + 1, 3).value
        next_notes = ws.cell(row + 1, 5).value
        kind_norm = normalize_header(next_kind)
        if kind_norm == "plant":
            activity = "PLANT"
        elif kind_norm == "harvest":
            activity = "HARVEST"
        else:
            continue
        events.append(
            {
                "section": "appendix",
                "year": event_date.year,
                "date": event_date,
                "activity": activity,
                "title": title,
                "notes": next_notes,
                "cell": f"B{row}",
                "activity_cell": f"C{row + 1}",
                "notes_cell": f"E{row + 1}",
            }
        )
    return events


def parse_csp3_events(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    if "CSP3" not in wb.sheetnames:
        raise CheckpointError(
            "site-history workbook has no CSP3 sheet",
            workbook=path.name,
        )
    ws = wb["CSP3"]
    events = _parse_structured(ws) + _parse_appendix(ws)
    wb.close()
    return events


def growing_periods(events: list[dict], year_min=2001, year_max=2017) -> list[dict]:
    by_year = defaultdict(lambda: {"plants": [], "harvests": []})
    for ev in events:
        year = ev.get("year")
        if year is None or year < year_min or year > year_max:
            continue
        if ev["date"] is None:
            continue
        if ev["activity"] == "PLANT":
            by_year[year]["plants"].append(ev)
        else:
            by_year[year]["harvests"].append(ev)

    rows = []
    for year in range(year_min, year_max + 1):
        plants = sorted(by_year[year]["plants"], key=lambda e: e["date"])
        harvests = sorted(by_year[year]["harvests"], key=lambda e: e["date"])
        skipped = []
        kept_plants = []
        for ev in plants:
            if should_skip_plant(ev.get("title"), ev.get("notes")):
                skipped.append(
                    {
                        "date": ev["date"].isoformat(),
                        "title": ev.get("title"),
                        "notes": ev.get("notes"),
                        "cell": ev.get("cell"),
                    }
                )
            else:
                kept_plants.append(ev)

        first_plant = kept_plants[0] if kept_plants else None
        last_harvest = harvests[-1] if harvests else None

        crop_texts = []
        for ev in (first_plant, last_harvest, *(kept_plants + harvests)):
            if ev:
                crop_texts.extend([ev.get("title"), ev.get("notes")])
        crop, crop_source = infer_crop(*crop_texts)

        plant_date = first_plant["date"] if first_plant else None
        harvest_date = last_harvest["date"] if last_harvest else None
        buffered_start = buffered_end = gee_end = None
        if plant_date and harvest_date:
            buffered_start, buffered_end = buffer_window(plant_date, harvest_date, BUFFER_DAYS)
            gee_end = gee_end_exclusive(buffered_end)

        rows.append(
            {
                "year": year,
                "crop": crop,
                "crop_source": crop_source,
                "window_type": "ops_plant_harvest",
                "ops_plant_date": plant_date.isoformat() if plant_date else None,
                "ops_harvest_date": harvest_date.isoformat() if harvest_date else None,
                "buffered_start": buffered_start.isoformat() if buffered_start else None,
                "buffered_end": buffered_end.isoformat() if buffered_end else None,
                "gee_filter_end_exclusive": gee_end.isoformat() if gee_end else None,
                "plant_title": first_plant.get("title") if first_plant else None,
                "plant_notes": first_plant.get("notes") if first_plant else None,
                "plant_cell": first_plant.get("cell") if first_plant else None,
                "harvest_title": last_harvest.get("title") if last_harvest else None,
                "harvest_notes": last_harvest.get("notes") if last_harvest else None,
                "harvest_cell": last_harvest.get("cell") if last_harvest else None,
                "n_plant_events": len(plants),
                "n_harvest_events": len(harvests),
                "skipped_plants": skipped,
                "workbook": SITE_HISTORY_FILE["name"],
                "sheet": SITE_HISTORY_FILE["sheet"],
            }
        )
    apply_rotation_fallback(rows)
    return rows


def apply_rotation_fallback(rows: list[dict]) -> None:
    """If plant/harvest notes have no crop token, use neighbors with an explicit flag.

    Never a silent odd/even assumption: both adjacent years must already be
    resolved to the same crop, and this year is labeled the other crop.
    """
    by_year = {r["year"]: r for r in rows}
    for row in rows:
        if row.get("crop"):
            continue
        prev = by_year.get(row["year"] - 1)
        nxt = by_year.get(row["year"] + 1)
        if not prev or not nxt:
            continue
        if not prev.get("crop") or not nxt.get("crop"):
            continue
        if prev["crop"] != nxt["crop"]:
            continue
        row["crop"] = "soybean" if prev["crop"] == "maize" else "maize"
        row["crop_source"] = "rotation_fallback_neighbors"


def apply_ops_checkpoints(rows: list[dict], workbook_name: str) -> None:
    by_year = {r["year"]: r for r in rows}
    for year, expected in OPS_CHECKPOINTS.items():
        row = by_year.get(year)
        if row is None:
            raise CheckpointError(
                f"ops checkpoint year {year} missing from site-history parse",
                workbook=workbook_name,
                sheet="CSP3",
            )
        got_plant = row["ops_plant_date"]
        exp_plant = expected["plant"].isoformat() if expected["plant"] else None
        if got_plant != exp_plant:
            raise CheckpointError(
                f"{year} plant date mismatch: expected {exp_plant}, got {got_plant}",
                workbook=workbook_name,
                sheet="CSP3",
                cell=row.get("plant_cell"),
                value=got_plant,
            )
        got_harv = row["ops_harvest_date"]
        exp_harv = expected["harvest"].isoformat() if expected["harvest"] else None
        if got_harv != exp_harv:
            raise CheckpointError(
                f"{year} harvest date mismatch: expected {exp_harv}, got {got_harv}",
                workbook=workbook_name,
                sheet="CSP3",
                cell=row.get("harvest_cell"),
                value=got_harv,
            )
        if row["crop"] != expected["crop"]:
            raise CheckpointError(
                f"{year} crop mismatch: expected {expected['crop']}, got {row['crop']} "
                f"(source={row['crop_source']})",
                workbook=workbook_name,
                sheet="CSP3",
                value=row["crop"],
            )


def extract_site_history(path: Path | None = None) -> list[dict]:
    path = path or (LOCAL_WORKBOOKS_DIR / SITE_HISTORY_FILE["name"])
    events = parse_csp3_events(path)
    rows = growing_periods(events)
    apply_ops_checkpoints(rows, path.name)
    return rows
