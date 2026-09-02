"""Maize phenology, All-IMZ MC_AGB, and HI from annual Site 3 biomass workbooks.

Dates come from V/R growth-stage sheets (never from blank Stage (DOY) HI templates).
End date is R6 sampling date, else the latest reproductive stage as provisional_end_date.
R5 is never labeled R6. Harvest dates are never substituted for R6.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from config import BUFFER_DAYS, PHENOLOGY_CHECKPOINTS, CheckpointError
from formulas import (
    buffer_window,
    coerce_moisture_fraction,
    gee_end_exclusive,
    harvest_index_direct,
    harvest_index_partition,
    mc_agb,
    planting_date_from_doy,
)
from headers import (
    as_float,
    cell_addr,
    find_header_column,
    find_label_value,
    header_matches,
    is_error_token,
    normalize_header,
)

STAGE_SHEET_RE = re.compile(r"^(V\d+|R\d+)\s*\((\d+)\)\s*$", re.I)
SKIP_SHEET_RE = re.compile(
    r"template|backup|means|pop data|contact|stage \(doy\)",
    re.I,
)


def _as_date(value) -> date | None:
    if value is None or is_error_token(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def parse_stage_sheet_name(name: str) -> tuple[str, str, int] | None:
    match = STAGE_SHEET_RE.match(str(name).strip())
    if not match:
        return None
    code = match.group(1).upper()
    doy = int(match.group(2))
    kind = "R" if code.startswith("R") else "V"
    return code, kind, doy


def list_stage_sheets(wb) -> list[dict]:
    out = []
    for name in wb.sheetnames:
        parsed = parse_stage_sheet_name(name)
        if parsed is None:
            continue
        code, kind, doy = parsed
        out.append({"name": name, "stage": code, "kind": kind, "sheet_doy": doy})
    return out


def read_planting_and_sampling(ws) -> dict:
    plant = find_label_value(ws, "doy of planting")
    stage = find_label_value(ws, "general growth stage")
    samp_date = find_label_value(ws, "sampling date")
    samp_doy = find_label_value(ws, "sampling doy")
    planting_doy = as_float(plant["value"]) if plant else None
    sampling_doy = as_float(samp_doy["value"]) if samp_doy else None
    growth_stage = None
    if stage and stage["value"] not in (None, "", "-") and not is_error_token(stage["value"]):
        growth_stage = str(stage["value"]).strip().upper()
    return {
        "a1": ws["A1"].value,
        "planting_doy": int(planting_doy) if planting_doy else None,
        "planting_doy_cell": plant["value_cell"] if plant else None,
        "planting_doy_header": plant["label"] if plant else None,
        "growth_stage": growth_stage,
        "growth_stage_cell": stage["value_cell"] if stage else None,
        "sampling_date": _as_date(samp_date["value"]) if samp_date else None,
        "sampling_date_cell": samp_date["value_cell"] if samp_date else None,
        "sampling_date_header": samp_date["label"] if samp_date else None,
        "sampling_doy": int(sampling_doy) if sampling_doy else None,
        "sampling_doy_cell": samp_doy["value_cell"] if samp_doy else None,
    }


def find_calculations_header_row(ws) -> int | None:
    for row, col, raw in _scan(ws, "pct moisture in fresh above ground biomass", "total dry above ground biomass"):
        if header_matches(raw, "pct moisture in fresh above ground biomass"):
            return row
        if header_matches(raw, "total dry above ground biomass") and row > 50:
            return row
    return None


def _scan(ws, *needles, max_row=None, max_col=None):
    from headers import find_cells

    return find_cells(ws, *needles, max_row=max_row, max_col=max_col)


def find_all_imz_mean_row(ws, header_row: int) -> int | None:
    for row in range(header_row + 1, (ws.max_row or header_row) + 1):
        b = ws.cell(row, 2).value
        c = ws.cell(row, 3).value
        if not isinstance(b, str):
            continue
        if "all imz" not in b.lower():
            continue
        c_norm = normalize_header(c)
        if c_norm == "mean":
            return row
    return None


def extract_mc_agb(ws) -> dict:
    header_row = find_calculations_header_row(ws)
    if header_row is None:
        return {"mc_agb": None, "reason": "no calculations header"}
    mean_row = find_all_imz_mean_row(ws, header_row)
    if mean_row is None:
        return {"mc_agb": None, "reason": "no All IMZ Mean row", "header_row": header_row}

    fresh_col = find_header_column(ws, header_row, "total fresh above ground biomass")
    dry_col = find_header_column(ws, header_row, "total dry above ground biomass")
    moist_col = find_header_column(ws, header_row, "pct moisture in fresh above ground biomass")
    repro_col = find_header_column(ws, header_row, "dry repro (kg/ha)")

    fresh = as_float(ws.cell(mean_row, fresh_col["col"]).value) if fresh_col else None
    dry = as_float(ws.cell(mean_row, dry_col["col"]).value) if dry_col else None
    moist_raw = as_float(ws.cell(mean_row, moist_col["col"]).value) if moist_col else None
    dry_repro = as_float(ws.cell(mean_row, repro_col["col"]).value) if repro_col else None

    reported = coerce_moisture_fraction(moist_raw) if moist_raw is not None else None
    ratio_of_means = None
    if fresh is not None and dry is not None:
        ratio_of_means = mc_agb(fresh, dry)
    # All-IMZ Mean of Pct. Moisture is the mean of plot-level (F-D)/F.
    # That is not equal to (mean F - mean D) / mean F; do not treat the
    # difference as a workbook error. Prefer the All-IMZ Mean cell.
    mc_value = reported if reported is not None else ratio_of_means

    return {
        "mc_agb": mc_value,
        "mc_agb_source": "all_imz_mean_pct_moisture" if reported is not None else "ratio_of_mean_masses",
        "mc_agb_ratio_of_means": ratio_of_means,
        "fresh_agb_kg_ha": fresh,
        "dry_agb_kg_ha": dry,
        "dry_repro_kg_ha": dry_repro,
        "pct_moisture_raw": moist_raw,
        "header_row": header_row,
        "all_imz_row": mean_row,
        "fresh_cell": cell_addr(mean_row, fresh_col["col"]) if fresh_col else None,
        "dry_cell": cell_addr(mean_row, dry_col["col"]) if dry_col else None,
        "mc_cell": cell_addr(mean_row, moist_col["col"]) if moist_col else None,
        "dry_repro_cell": cell_addr(mean_row, repro_col["col"]) if repro_col else None,
        "fresh_header": fresh_col["header"] if fresh_col else None,
        "dry_header": dry_col["header"] if dry_col else None,
        "mc_header": moist_col["header"] if moist_col else None,
    }


def extract_direct_hi(ws) -> dict | None:
    """Read Harvest Index from a Stage (DOY) HI sheet when Dry Grain is numeric."""
    header_row = find_calculations_header_row(ws)
    if header_row is None:
        return None
    grain_in = None
    for row, col, raw in _scan(ws, "dry grain (g)", max_row=20, max_col=20):
        grain_in = (row, col, raw)
        break
    if grain_in is None:
        return None

    numeric_grain = 0
    for row in range(8, 60):
        if as_float(ws.cell(row, grain_in[1]).value) is not None:
            numeric_grain += 1
    if numeric_grain == 0:
        return {
            "hi": None,
            "hi_method": None,
            "hi_status": "direct_grain_inputs_missing",
            "sheet": ws.title,
        }

    mean_row = find_all_imz_mean_row(ws, header_row)
    hi_col = find_header_column(ws, header_row, "harvest index")
    grain_col = find_header_column(ws, header_row, "dry grain (kg/ha)")
    dry_col = find_header_column(ws, header_row, "total dry above ground biomass")
    if mean_row is None or hi_col is None or grain_col is None or dry_col is None:
        return None
    hi_val = as_float(ws.cell(mean_row, hi_col["col"]).value)
    grain = as_float(ws.cell(mean_row, grain_col["col"]).value)
    dry = as_float(ws.cell(mean_row, dry_col["col"]).value)
    if hi_val is None or grain is None or dry is None:
        return {
            "hi": None,
            "hi_method": None,
            "hi_status": "direct_grain_all_imz_nonnumeric",
            "sheet": ws.title,
        }
    computed = harvest_index_direct(grain, dry)
    if abs(computed - hi_val) > 1e-4:
        raise CheckpointError(
            f"direct HI cell disagrees with dry grain / dry AGB: cell={hi_val}, computed={computed}",
            sheet=ws.title,
            cell=cell_addr(mean_row, hi_col["col"]),
            value=hi_val,
        )
    return {
        "hi": computed,
        "hi_method": "direct_dry_grain",
        "hi_status": "ok",
        "dry_grain_kg_ha": grain,
        "dry_agb_kg_ha": dry,
        "sheet": ws.title,
        "hi_cell": cell_addr(mean_row, hi_col["col"]),
        "hi_header": hi_col["header"],
    }


def extract_partition_hi(ws, dry_repro, dry_agb) -> dict:
    """Kernel / subsample weights from Reproductive Sample Partitioning."""
    header_row = None
    weight_col = None
    for row, col, raw in _scan(ws, "weight (grams)", "reproductive sample partitioning"):
        if header_matches(raw, "weight (grams)"):
            header_row = row
            weight_col = col
            break
    if header_row is None or weight_col is None:
        return {
            "hi": None,
            "hi_method": None,
            "hi_status": "partition_block_missing",
        }

    subsample = kernel = None
    subsample_cell = kernel_cell = None
    subsample_header = kernel_header = None
    for row in range(header_row, min(header_row + 12, (ws.max_row or header_row) + 1)):
        labels = [ws.cell(row, c).value for c in range(1, weight_col)]
        blob = " ".join(str(x) for x in labels if x is not None)
        blob_n = normalize_header(blob)
        val = as_float(ws.cell(row, weight_col).value)
        if "reproductive subsample" in blob_n:
            subsample = val
            subsample_cell = cell_addr(row, weight_col)
            subsample_header = blob
        elif re.search(r"^kernel\s*$", blob_n) or blob_n.endswith(" kernel") or blob_n == "kernel":
            # Ignore "Kernel Weight (mg per kernel)"
            if "mg per kernel" in blob_n:
                continue
            kernel = val
            kernel_cell = cell_addr(row, weight_col)
            kernel_header = blob

    if subsample is None or kernel is None:
        return {
            "hi": None,
            "hi_method": None,
            "hi_status": "kernel_or_subsample_nonnumeric",
            "kernel_g": kernel,
            "subsample_g": subsample,
            "weight_header_cell": cell_addr(header_row, weight_col),
            "sheet": ws.title,
        }
    if dry_repro is None or dry_agb is None:
        return {
            "hi": None,
            "hi_method": None,
            "hi_status": "missing_dry_repro_or_agb",
            "kernel_g": kernel,
            "subsample_g": subsample,
            "sheet": ws.title,
        }
    part = harvest_index_partition(dry_repro, dry_agb, kernel, subsample)
    return {
        "hi": part["hi"],
        "hi_method": "partition-derived",
        "hi_status": "ok",
        "kernel_fraction": part["kernel_fraction"],
        "dry_grain_kg_ha": part["dry_grain_kg_ha"],
        "kernel_g": kernel,
        "subsample_g": subsample,
        "kernel_cell": kernel_cell,
        "subsample_cell": subsample_cell,
        "kernel_header": kernel_header,
        "subsample_header": subsample_header,
        "weight_header": ws.cell(header_row, weight_col).value,
        "weight_header_cell": cell_addr(header_row, weight_col),
        "sheet": ws.title,
    }


def choose_end_stage(stage_sheets: list[dict], meta_by_name: dict) -> dict:
    r_sheets = [s for s in stage_sheets if s["kind"] == "R"]
    if not r_sheets:
        raise CheckpointError("no reproductive-stage sheets found")
    r6 = [s for s in r_sheets if s["stage"] == "R6"]
    if r6:
        chosen = max(r6, key=lambda s: s["sheet_doy"])
        provisional = False
    else:
        chosen = max(r_sheets, key=lambda s: s["sheet_doy"])
        provisional = True
        if chosen["stage"] == "R6":
            raise CheckpointError("internal error: R5 selected as R6")
    meta = meta_by_name[chosen["name"]]
    stage_label = (meta.get("growth_stage") or chosen["stage"]).upper()
    if provisional and stage_label == "R6":
        raise CheckpointError(
            "refusing to label a non-R6 sheet as R6",
            sheet=chosen["name"],
            value=stage_label,
        )
    if not provisional and stage_label not in {"R6", None}:
        # R6 sheet should say R6
        if stage_label != "R6":
            raise CheckpointError(
                f"R6 sheet growth stage is {stage_label!r}, expected R6",
                sheet=chosen["name"],
                cell=meta.get("growth_stage_cell"),
                value=stage_label,
            )
    return {
        "sheet": chosen["name"],
        "end_stage": chosen["stage"] if provisional else "R6",
        "end_is_provisional": provisional,
        "sheet_doy": chosen["sheet_doy"],
        "meta": meta,
    }


def extract_maize_year(path: Path, year: int) -> dict:
    wb = load_workbook(path, data_only=True)
    try:
        a1 = None
        for name in wb.sheetnames:
            parsed = parse_stage_sheet_name(name)
            if parsed:
                a1 = wb[name]["A1"].value
                break
        a1_text = str(a1) if a1 is not None else ""
        if "soy" in a1_text.lower():
            raise CheckpointError(
                f"{path.name} A1 is soybean; maize extractor must not run",
                workbook=path.name,
                sheet=name,
                cell="A1",
                value=a1,
            )
        if "maize" not in a1_text.lower() and "corn" not in a1_text.lower():
            raise CheckpointError(
                f"{path.name} A1 is not explicit maize",
                workbook=path.name,
                cell="A1",
                value=a1,
            )

        stage_sheets = list_stage_sheets(wb)
        meta_by_name = {}
        planting_doys = set()
        for spec in stage_sheets:
            ws = wb[spec["name"]]
            meta = read_planting_and_sampling(ws)
            meta_by_name[spec["name"]] = meta
            if meta["planting_doy"]:
                planting_doys.add(meta["planting_doy"])

        if not planting_doys:
            raise CheckpointError("no planting DOY on V/R stage sheets", workbook=path.name)
        if len(planting_doys) != 1:
            raise CheckpointError(
                f"planting DOY disagrees across V/R sheets: {sorted(planting_doys)}",
                workbook=path.name,
            )

        end = choose_end_stage(stage_sheets, meta_by_name)
        end_ws = wb[end["sheet"]]
        meta = end["meta"]
        planting_doy = meta["planting_doy"] or next(iter(planting_doys))
        planting_date = planting_date_from_doy(year, planting_doy)
        sampling_date = meta["sampling_date"]
        if sampling_date is None and meta["sampling_doy"]:
            sampling_date = planting_date_from_doy(year, meta["sampling_doy"])
        if sampling_date is None:
            raise CheckpointError(
                "end-stage sheet has no sampling date",
                workbook=path.name,
                sheet=end["sheet"],
            )

        buffered_start, buffered_end = buffer_window(planting_date, sampling_date, BUFFER_DAYS)
        mc = extract_mc_agb(end_ws)

        hi = {
            "hi": None,
            "hi_method": None,
            "hi_status": "not_computed",
        }
        if "Stage (DOY) HI" in wb.sheetnames:
            direct = extract_direct_hi(wb["Stage (DOY) HI"])
            if direct and direct.get("hi") is not None:
                hi = direct
            elif direct:
                hi = {**hi, **direct}

        if hi.get("hi") is None:
            part = extract_partition_hi(end_ws, mc.get("dry_repro_kg_ha"), mc.get("dry_agb_kg_ha"))
            if end["end_is_provisional"] and part.get("hi") is not None:
                part["hi_status"] = "provisional_r5"
            hi = part

        return {
            "year": year,
            "crop": "maize",
            "a1_raw": a1_text,
            "workbook": path.name,
            "window_type": "phenology_rstage",
            "planting_doy": planting_doy,
            "planting_date": planting_date.isoformat(),
            "planting_doy_cell": meta.get("planting_doy_cell"),
            "planting_doy_header": meta.get("planting_doy_header"),
            "end_stage": end["end_stage"],
            "end_is_provisional": end["end_is_provisional"],
            "end_sheet": end["sheet"],
            "sampling_date": sampling_date.isoformat(),
            "sampling_doy": meta.get("sampling_doy") or end["sheet_doy"],
            "sampling_date_cell": meta.get("sampling_date_cell"),
            "sampling_date_header": meta.get("sampling_date_header"),
            "growth_stage_cell": meta.get("growth_stage_cell"),
            "buffered_start": buffered_start.isoformat(),
            "buffered_end": buffered_end.isoformat(),
            "gee_filter_end_exclusive": gee_end_exclusive(buffered_end).isoformat(),
            "mc_agb": mc.get("mc_agb"),
            "mc_agb_source": mc.get("mc_agb_source"),
            "fresh_agb_kg_ha": mc.get("fresh_agb_kg_ha"),
            "dry_agb_kg_ha": mc.get("dry_agb_kg_ha"),
            "dry_repro_kg_ha": mc.get("dry_repro_kg_ha"),
            "mc_cell": mc.get("mc_cell"),
            "mc_header": mc.get("mc_header"),
            "fresh_cell": mc.get("fresh_cell"),
            "dry_cell": mc.get("dry_cell"),
            "dry_repro_cell": mc.get("dry_repro_cell"),
            "all_imz_row": mc.get("all_imz_row"),
            **hi,
        }
    finally:
        wb.close()


def _close(a, b, tol) -> bool:
    if a is None and b is None:
        return True
    if a is None or b is None:
        return False
    return abs(float(a) - float(b)) <= tol


def apply_phenology_checkpoints(row: dict) -> None:
    year = row["year"]
    exp = PHENOLOGY_CHECKPOINTS.get(year)
    if not exp:
        return
    checks = [
        ("planting_doy", row["planting_doy"], exp["planting_doy"], 0, row.get("planting_doy_cell")),
        ("end_stage", row["end_stage"], exp["end_stage"], None, row.get("end_sheet")),
        ("end_is_provisional", row["end_is_provisional"], exp["end_is_provisional"], None, None),
        ("sampling_doy", row["sampling_doy"], exp["sampling_doy"], 0, row.get("sampling_date_cell")),
        (
            "buffered_start",
            date.fromisoformat(row["buffered_start"]),
            exp["buffered_start"],
            None,
            None,
        ),
        (
            "buffered_end",
            date.fromisoformat(row["buffered_end"]),
            exp["buffered_end"],
            None,
            None,
        ),
        ("mc_agb", row.get("mc_agb"), exp["mc_agb"], 1e-8, row.get("mc_cell")),
        ("hi", row.get("hi"), exp["hi"], 1e-6, row.get("hi_cell") or row.get("kernel_cell")),
    ]
    for name, got, expected, tol, cell in checks:
        if name in {"planting_doy", "sampling_doy", "end_stage", "end_is_provisional", "buffered_start", "buffered_end"}:
            if got != expected:
                raise CheckpointError(
                    f"{year} {name} mismatch: expected {expected}, got {got}",
                    workbook=row.get("workbook"),
                    sheet=row.get("end_sheet"),
                    cell=cell,
                    value=got,
                )
        else:
            if expected is None:
                if got is not None:
                    raise CheckpointError(
                        f"{year} {name} expected NA, got {got}",
                        workbook=row.get("workbook"),
                        sheet=row.get("end_sheet"),
                        cell=cell,
                        value=got,
                    )
            elif not _close(got, expected, tol):
                raise CheckpointError(
                    f"{year} {name} mismatch: expected {expected}, got {got}",
                    workbook=row.get("workbook"),
                    sheet=row.get("end_sheet"),
                    cell=cell,
                    value=got,
                )
    if exp.get("hi_method") and row.get("hi") is not None:
        if row.get("hi_method") != exp["hi_method"]:
            raise CheckpointError(
                f"{year} HI method mismatch: expected {exp['hi_method']}, got {row.get('hi_method')}",
                workbook=row.get("workbook"),
                value=row.get("hi_method"),
            )


def extract_all_maize(workbook_by_year: dict[int, dict]) -> list[dict]:
    rows = []
    for year, rec in sorted(workbook_by_year.items()):
        if rec.get("a1_crop") != "maize":
            continue
        row = extract_maize_year(Path(rec["path"]), year)
        apply_phenology_checkpoints(row)
        rows.append(row)
    return rows
