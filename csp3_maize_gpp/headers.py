"""Normalized header matching. Never rely on column letters alone."""

from __future__ import annotations

import re
from typing import Any

from openpyxl.utils import get_column_letter

ERROR_TOKENS = {"#DIV/0!", "#VALUE!", "#REF!", "#N/A", "#NAME?", "#NULL!", "#NUM!"}


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ")
    text = text.replace("\xa0", " ").replace("\u200b", "")
    text = text.strip().lower()
    text = re.sub(r"\s+", " ", text)
    text = text.replace(".", "")
    return text


def is_error_token(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str) and value.strip() in ERROR_TOKENS:
        return True
    text = str(value).strip()
    return text in ERROR_TOKENS


def is_numeric(value: Any) -> bool:
    if value is None or is_error_token(value) or value == "-":
        return False
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    try:
        float(str(value).replace(",", ""))
        return True
    except (TypeError, ValueError):
        return False


def as_float(value: Any) -> float | None:
    if not is_numeric(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).replace(",", ""))


def cell_addr(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def header_matches(cell_value: Any, *needles: str) -> bool:
    norm = normalize_header(cell_value)
    if not norm:
        return False
    for needle in needles:
        n = normalize_header(needle)
        if n and n in norm:
            return True
    return False


def find_cells(ws, *needles: str, max_row: int | None = None, max_col: int | None = None):
    """Yield (row, col, raw_header) for cells whose normalized text contains a needle."""
    max_row = max_row or ws.max_row or 1
    max_col = max_col or ws.max_column or 1
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            value = ws.cell(row, col).value
            if header_matches(value, *needles):
                yield row, col, value


def value_right_of(ws, row: int, col: int, max_look: int = 8) -> tuple[Any, str | None]:
    """Return the first non-empty cell to the right of a label."""
    for offset in range(1, max_look + 1):
        cell = ws.cell(row, col + offset)
        if cell.value is not None and str(cell.value).strip() != "":
            return cell.value, cell_addr(row, col + offset)
    return None, None


def find_label_value(ws, *needles: str, search_rows: int = 12, search_cols: int = 20):
    """Find a label in the header block and the first value to its right."""
    for row, col, raw in find_cells(ws, *needles, max_row=search_rows, max_col=search_cols):
        value, addr = value_right_of(ws, row, col)
        return {
            "label": raw,
            "label_cell": cell_addr(row, col),
            "value": value,
            "value_cell": addr,
            "row": row,
            "col": col,
        }
    return None


def find_header_column(ws, header_row: int, *needles: str) -> dict | None:
    max_col = ws.max_column or 1
    for col in range(1, max_col + 1):
        raw = ws.cell(header_row, col).value
        if header_matches(raw, *needles):
            return {
                "header": raw,
                "header_cell": cell_addr(header_row, col),
                "col": col,
                "row": header_row,
            }
    return None
